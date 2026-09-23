from __future__ import annotations

from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Image,
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data" / "processed"
MODEL = ROOT / "model"
FIGURES = ROOT / "figures"
REPORT = ROOT / "report"

NAVY = "16213E"
RED = "E53935"
BLUE = "2F6BFF"
LIGHT_BLUE = "EAF0FF"
LIGHT_GRAY = "F3F5F7"
MID_GRAY = "D7DCE2"
GREEN = "1F9D72"
ORANGE = "F59E0B"

SHARES_M = 7.533015
BEGINNING_EQUITY_BN = 1548.3
NET_CASH_BN = 76.307
WACC = 0.085
TERMINAL_GROWTH = 0.025
COE = 0.09
PAYOUT = 0.15
TERMINAL_ROE = 0.14
REFERENCE_PRICE = 1_300_000
PER_MULTIPLE = 14.72


def dcf_value(fcff: list[float], wacc: float, growth: float) -> tuple[float, float, float]:
    pv_stage1 = sum(value / (1 + wacc) ** (year + 0.5) for year, value in enumerate(fcff))
    terminal_value = fcff[-1] * (1 + growth) / (wacc - growth)
    pv_terminal = terminal_value / (1 + wacc) ** (len(fcff) - 0.5)
    equity_value = pv_stage1 + pv_terminal + NET_CASH_BN
    return equity_value, pv_stage1, pv_terminal


def rim_value(net_income: list[float], coe: float, terminal_roe: float) -> tuple[float, list[dict]]:
    equity = BEGINNING_EQUITY_BN
    pv_residual = 0.0
    rows = []
    for year, income in zip(range(2026, 2031), net_income):
        residual = income - coe * equity
        pv = residual / (1 + coe) ** (year - 2026 + 0.5)
        dividend = income * PAYOUT
        ending_equity = equity + income - dividend
        rows.append(
            {
                "year": year,
                "beginning_equity": equity,
                "net_income": income,
                "dividend": dividend,
                "ending_equity": ending_equity,
                "residual_income": residual,
                "pv_residual_income": pv,
            }
        )
        pv_residual += pv
        equity = ending_equity
    terminal_ri = (terminal_roe - coe) * equity
    terminal_value = terminal_ri / (coe - TERMINAL_GROWTH)
    pv_terminal = terminal_value / (1 + coe) ** 4.5
    return BEGINNING_EQUITY_BN + pv_residual + pv_terminal, rows


def style_sheet(ws, widths: dict[int, float]) -> None:
    thin = Side(style="thin", color=MID_GRAY)
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="center")
            if cell.value is not None:
                cell.border = Border(bottom=thin)


def title_row(ws, title: str, end_col: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    cell = ws.cell(1, 1, title)
    cell.font = Font(size=18, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 30


def header(ws, row: int, start_col: int, values: list[str]) -> None:
    for i, value in enumerate(values, start_col):
        cell = ws.cell(row, i, value)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center")


def build_dcf_workbook(base: pd.DataFrame) -> Path:
    path = MODEL / "DCF_model.xlsx"
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    inputs = wb.create_sheet("Inputs")
    forecast = wb.create_sheet("Forecast")
    sensitivity = wb.create_sheet("Sensitivity")

    fcff = base["fcff_krw_bn"].tolist()
    equity_value, pv_stage1, pv_terminal = dcf_value(fcff, WACC, TERMINAL_GROWTH)
    price = equity_value * 1000 / SHARES_M

    title_row(summary, "Samyang Foods DCF Valuation", 5)
    summary.append([])
    summary.append(["Valuation date", "2026-06-30", None, "Reference price", REFERENCE_PRICE])
    summary.append(["DCF equity value (KRW bn)", equity_value, None, "DCF value/share", price])
    summary.append(["PV of explicit FCFF", pv_stage1, None, "Upside vs reference", price / REFERENCE_PRICE - 1])
    summary.append(["PV of terminal value", pv_terminal, None, "Status", "2026H1 filing anchored"])
    summary.append(["Net cash", NET_CASH_BN, None, "Model convention", "Mid-year discounting"])
    summary["B4"].number_format = '#,##0.0'
    summary["B5"].number_format = '#,##0.0'
    summary["B6"].number_format = '#,##0.0'
    summary["B7"].number_format = '#,##0.0'
    summary["E3"].number_format = '#,##0'
    summary["E4"].number_format = '#,##0'
    summary["E5"].number_format = '0.0%'
    summary.freeze_panes = "A3"
    style_sheet(summary, {1: 30, 2: 20, 3: 4, 4: 24, 5: 28})

    title_row(inputs, "DCF Inputs", 6)
    header(inputs, 3, 1, ["Category", "Item", "Value", "Unit", "Status", "Source / note"])
    input_rows = [
        ["Common", "Shares outstanding", SHARES_M, "mn shares", "Filed", "7,533,015 common shares"],
        ["Common", "Net cash", NET_CASH_BN, "KRW bn", "Filed", "2026H1 cash + ST financial assets - debt - leases"],
        ["DCF", "WACC", WACC, "%", "Assumption", "Independent model assumption"],
        ["DCF", "Terminal growth", TERMINAL_GROWTH, "%", "Assumption", "Long-run nominal growth"],
        ["DCF", "Discount timing", "Mid-year", "text", "Assumption", "2026E FCFF discounted by 0.5 year"],
        ["Market", "Reference price", REFERENCE_PRICE, "KRW/share", "Observed", "2026-09-18 close"],
    ]
    for row in input_rows:
        inputs.append(row)
    for row in range(4, 10):
        if inputs.cell(row, 2).value in {"WACC", "Terminal growth"}:
            inputs.cell(row, 3).number_format = "0.0%"
    style_sheet(inputs, {1: 14, 2: 24, 3: 18, 4: 16, 5: 16, 6: 52})
    inputs.freeze_panes = "A4"

    title_row(forecast, "DCF Forecast & Discounting", 14)
    headers = [
        "Year", "Revenue", "Growth", "EBIT margin", "EBIT", "Tax", "NOPAT",
        "D&A", "CAPEX", "ΔNWC", "FCFF", "Discount factor", "PV FCFF", "Unit",
    ]
    header(forecast, 3, 1, headers)
    for i, row in enumerate(base.itertuples(index=False), 4):
        forecast.cell(i, 1, row.year)
        forecast.cell(i, 2, row.revenue_krw_bn)
        forecast.cell(i, 3, row.revenue_growth)
        forecast.cell(i, 4, row.operating_margin)
        forecast.cell(i, 5, row.ebit_krw_bn)
        forecast.cell(i, 6, row.tax_rate)
        forecast.cell(i, 7, f"=E{i}*(1-F{i})")
        forecast.cell(i, 8, row.d_and_a_krw_bn)
        forecast.cell(i, 9, row.capex_krw_bn)
        forecast.cell(i, 10, row.delta_core_nwc_krw_bn)
        forecast.cell(i, 11, f"=G{i}+H{i}-I{i}-J{i}")
        forecast.cell(i, 12, f"=1/(1+Inputs!$C$6)^(A{i}-2026+0.5)")
        forecast.cell(i, 13, f"=K{i}*L{i}")
        forecast.cell(i, 14, "KRW bn")
        for col in [3, 4, 6, 12]:
            forecast.cell(i, col).number_format = "0.0%"
        for col in [2, 5, 7, 8, 9, 10, 11, 13]:
            forecast.cell(i, col).number_format = '#,##0.0'
    style_sheet(forecast, {1: 10, 2: 14, 3: 12, 4: 14, 5: 14, 6: 10, 7: 14, 8: 12, 9: 12, 10: 12, 11: 14, 12: 16, 13: 14, 14: 12})
    forecast.freeze_panes = "A4"

    title_row(sensitivity, "DCF Sensitivity — Value per Share (KRW)", 5)
    growths = [0.020, 0.025, 0.030]
    waccs = [0.075, 0.080, 0.085, 0.090, 0.095]
    header(sensitivity, 3, 1, ["WACC / g"] + [f"{g:.1%}" for g in growths])
    for r, wacc in enumerate(waccs, 4):
        sensitivity.cell(r, 1, wacc)
        sensitivity.cell(r, 1).number_format = "0.0%"
        for c, growth in enumerate(growths, 2):
            value, _, _ = dcf_value(fcff, wacc, growth)
            sensitivity.cell(r, c, value * 1000 / SHARES_M)
            sensitivity.cell(r, c).number_format = '#,##0'
    sensitivity.conditional_formatting.add(
        "B4:D8",
        ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50, mid_color="FFEB84", end_type="max", end_color="63BE7B"),
    )
    style_sheet(sensitivity, {1: 16, 2: 18, 3: 18, 4: 18, 5: 4})

    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        ws.auto_filter.ref = ws.dimensions
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.save(path)
    return path


def build_rim_workbook(base: pd.DataFrame) -> Path:
    path = MODEL / "RIM_model.xlsx"
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    inputs = wb.create_sheet("Inputs")
    forecast = wb.create_sheet("Residual Income")
    sensitivity = wb.create_sheet("Sensitivity")

    equity_value, rows = rim_value(base["net_income_krw_bn"].tolist(), COE, TERMINAL_ROE)
    price = equity_value * 1000 / SHARES_M

    title_row(summary, "Samyang Foods Residual Income Model", 5)
    summary.append([])
    summary.append(["Valuation date", "2026-06-30", None, "Reference price", REFERENCE_PRICE])
    summary.append(["Beginning common equity", BEGINNING_EQUITY_BN, None, "RIM value/share", price])
    summary.append(["RIM equity value (KRW bn)", equity_value, None, "Upside vs reference", price / REFERENCE_PRICE - 1])
    summary.append(["Cost of equity", COE, None, "Terminal ROE", TERMINAL_ROE])
    summary.append(["Terminal growth", TERMINAL_GROWTH, None, "Payout ratio", PAYOUT])
    for cell in ["B4", "B5"]:
        summary[cell].number_format = '#,##0.0'
    for cell in ["B6", "B7", "E5", "E6", "E7"]:
        summary[cell].number_format = "0.0%"
    summary["E3"].number_format = '#,##0'
    summary["E4"].number_format = '#,##0'
    style_sheet(summary, {1: 30, 2: 20, 3: 4, 4: 24, 5: 22})

    title_row(inputs, "RIM Inputs", 6)
    header(inputs, 3, 1, ["Category", "Item", "Value", "Unit", "Status", "Source / note"])
    input_rows = [
        ["Common", "Beginning common equity", BEGINNING_EQUITY_BN, "KRW bn", "Filed", "2026H1 DART filing; consolidated equity"],
        ["Common", "Shares outstanding", SHARES_M, "mn shares", "Filed", "7,533,015 common shares"],
        ["RIM", "Cost of equity", COE, "%", "Assumption", "Independent model assumption"],
        ["RIM", "Payout ratio", PAYOUT, "%", "Assumption", "Applied to Base net income"],
        ["RIM", "Terminal ROE", TERMINAL_ROE, "%", "Assumption", "Mature-state ROE"],
        ["RIM", "Terminal growth", TERMINAL_GROWTH, "%", "Assumption", "Long-run nominal growth"],
        ["Market", "Reference price", REFERENCE_PRICE, "KRW/share", "Observed", "2026-09-18 close"],
    ]
    for row in input_rows:
        inputs.append(row)
    for row in range(4, 11):
        if inputs.cell(row, 2).value in {"Cost of equity", "Payout ratio", "Terminal ROE", "Terminal growth"}:
            inputs.cell(row, 3).number_format = "0.0%"
    style_sheet(inputs, {1: 14, 2: 28, 3: 18, 4: 16, 5: 16, 6: 52})
    inputs.freeze_panes = "A4"

    title_row(forecast, "Residual Income Forecast", 9)
    header(forecast, 3, 1, ["Year", "Beginning equity", "Net income", "Dividend", "Ending equity", "Equity charge", "Residual income", "PV factor", "PV residual income"])
    for i, row in enumerate(rows, 4):
        forecast.cell(i, 1, row["year"])
        forecast.cell(i, 2, row["beginning_equity"])
        forecast.cell(i, 3, row["net_income"])
        forecast.cell(i, 4, f"=C{i}*Inputs!$C$7")
        forecast.cell(i, 5, f"=B{i}+C{i}-D{i}")
        forecast.cell(i, 6, f"=B{i}*Inputs!$C$6")
        forecast.cell(i, 7, f"=C{i}-F{i}")
        forecast.cell(i, 8, f"=1/(1+Inputs!$C$6)^(A{i}-2026+0.5)")
        forecast.cell(i, 9, f"=G{i}*H{i}")
        for col in range(2, 8):
            forecast.cell(i, col).number_format = '#,##0.0'
        forecast.cell(i, 8).number_format = "0.000"
        forecast.cell(i, 9).number_format = '#,##0.0'
    style_sheet(forecast, {1: 10, 2: 18, 3: 15, 4: 14, 5: 17, 6: 16, 7: 17, 8: 13, 9: 20})
    forecast.freeze_panes = "A4"

    title_row(sensitivity, "RIM Sensitivity — Value per Share (KRW)", 6)
    coes = [0.080, 0.085, 0.090, 0.095, 0.100]
    roes = [0.12, 0.14, 0.16]
    header(sensitivity, 3, 1, ["COE / terminal ROE"] + [f"{x:.0%}" for x in roes])
    for r, coe in enumerate(coes, 4):
        sensitivity.cell(r, 1, coe)
        sensitivity.cell(r, 1).number_format = "0.0%"
        for c, roe in enumerate(roes, 2):
            value, _ = rim_value(base["net_income_krw_bn"].tolist(), coe, roe)
            sensitivity.cell(r, c, value * 1000 / SHARES_M)
            sensitivity.cell(r, c).number_format = '#,##0'
    sensitivity.conditional_formatting.add(
        "B4:D8",
        ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50, mid_color="FFEB84", end_type="max", end_color="63BE7B"),
    )
    style_sheet(sensitivity, {1: 23, 2: 18, 3: 18, 4: 18, 5: 4})

    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        ws.auto_filter.ref = ws.dimensions
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.save(path)
    return path


def build_figures(hist: pd.DataFrame, forecast: pd.DataFrame, valuation: dict[str, float]) -> list[Path]:
    FIGURES.mkdir(exist_ok=True)
    plt.style.use("seaborn-v0_8-whitegrid")

    fig, ax1 = plt.subplots(figsize=(9, 4.8))
    ax1.bar(hist["year"], hist["revenue_krw_mn"] / 1000, color=f"#{BLUE}", alpha=0.85, label="Revenue")
    ax1.set_ylabel("Revenue (KRW bn)")
    ax1.set_title("Historical Revenue and Operating Margin")
    ax2 = ax1.twinx()
    margin = hist["operating_profit_krw_mn"] / hist["revenue_krw_mn"]
    ax2.plot(hist["year"], margin * 100, color=f"#{RED}", marker="o", linewidth=2.5, label="Operating margin")
    ax2.set_ylabel("Operating margin (%)")
    ax1.set_xticks(hist["year"])
    fig.tight_layout()
    p1 = FIGURES / "historical_financials.png"
    fig.savefig(p1, dpi=180, bbox_inches="tight")
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(9, 4.8))
    for scenario, group in forecast.groupby("scenario"):
        color = f"#{ {'Base': BLUE, 'Bull': GREEN, 'Bear': ORANGE}[scenario] }"
        ax.plot(group["year"], group["revenue_krw_bn"], marker="o", linewidth=2.4, label=scenario, color=color)
    ax.set_title("Revenue Scenarios (2026E–2030E)")
    ax.set_ylabel("KRW bn")
    ax.set_xticks(sorted(forecast["year"].unique()))
    ax.legend(frameon=False)
    fig.tight_layout()
    p2 = FIGURES / "revenue_scenarios.png"
    fig.savefig(p2, dpi=180, bbox_inches="tight")
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(8.5, 4.8))
    names = ["Reference price", "DCF", "PER", "RIM"]
    values = [REFERENCE_PRICE, valuation["DCF"], valuation["PER"], valuation["RIM"]]
    bars = ax.bar(names, np.array(values) / 1_000_000, color=["#6B7280", f"#{BLUE}", f"#{GREEN}", f"#{ORANGE}"])
    ax.set_ylabel("KRW million per share")
    ax.set_title("Valuation Cross-check")
    for bar, value in zip(bars, values):
        ax.text(bar.get_x() + bar.get_width() / 2, bar.get_height() + 0.025, f"{value/1_000_000:.2f}", ha="center", fontsize=9)
    ax.set_ylim(0, max(np.array(values) / 1_000_000) * 1.2)
    fig.tight_layout()
    p3 = FIGURES / "valuation_crosscheck.png"
    fig.savefig(p3, dpi=180, bbox_inches="tight")
    plt.close(fig)
    return [p1, p2, p3]


def build_pdf(hist: pd.DataFrame, forecast: pd.DataFrame, valuation: dict[str, float], figures: list[Path]) -> Path:
    path = REPORT / "Samyang_Foods_Equity_Research_2026H1.pdf"
    pdfmetrics.registerFont(TTFont("NanumGothic", ROOT / "assets" / "fonts" / "NanumGothic-Regular.ttf"))
    pdfmetrics.registerFont(TTFont("NanumGothic-Bold", ROOT / "assets" / "fonts" / "NanumGothic-Bold.ttf"))
    regular = "NanumGothic"
    bold = "NanumGothic-Bold"

    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="KTitle", fontName=bold, fontSize=27, leading=35, textColor=colors.HexColor(f"#{NAVY}"), alignment=TA_CENTER, spaceAfter=12))
    styles.add(ParagraphStyle(name="KSubtitle", fontName=regular, fontSize=12, leading=19, textColor=colors.HexColor("#475569"), alignment=TA_CENTER))
    styles.add(ParagraphStyle(name="KH1", fontName=bold, fontSize=18, leading=25, textColor=colors.HexColor(f"#{NAVY}"), spaceAfter=12))
    styles.add(ParagraphStyle(name="KH2", fontName=bold, fontSize=12, leading=18, textColor=colors.HexColor(f"#{BLUE}"), spaceBefore=8, spaceAfter=6))
    styles.add(ParagraphStyle(name="KBody", fontName=regular, fontSize=9.4, leading=15, textColor=colors.HexColor("#273043"), spaceAfter=7))
    styles.add(ParagraphStyle(name="KSmall", fontName=regular, fontSize=7.5, leading=11, textColor=colors.HexColor("#4B5563"), spaceAfter=4))
    styles.add(ParagraphStyle(name="KCallout", fontName=bold, fontSize=10.5, leading=16, textColor=colors.HexColor(f"#{NAVY}"), backColor=colors.HexColor(f"#{LIGHT_BLUE}"), borderPadding=9, spaceAfter=8))
    styles.add(ParagraphStyle(name="KSource", fontName=regular, fontSize=7.2, leading=11, textColor=colors.HexColor("#475569"), leftIndent=8, firstLineIndent=-8, spaceAfter=4))

    def p(text: str, style: str = "KBody") -> Paragraph:
        return Paragraph(text, styles[style])

    def table(data, widths=None, font_size=8.1):
        t = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
        t.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{NAVY}")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), bold),
                    ("FONTNAME", (0, 1), (-1, -1), regular),
                    ("FONTSIZE", (0, 0), (-1, -1), font_size),
                    ("LEADING", (0, 0), (-1, -1), font_size + 3),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor(f"#{MID_GRAY}")),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor(f"#{LIGHT_GRAY}")]),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 5),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                    ("TOPPADDING", (0, 0), (-1, -1), 5),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ]
            )
        )
        return t

    def page_number(canvas, doc):
        canvas.saveState()
        canvas.setFont(regular, 7)
        canvas.setFillColor(colors.HexColor("#6B7280"))
        canvas.drawString(18 * mm, 12 * mm, "Samyang Foods Equity Research · 2026H1")
        canvas.drawRightString(A4[0] - 18 * mm, 12 * mm, str(doc.page))
        canvas.restoreState()

    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=17 * mm, rightMargin=17 * mm, topMargin=16 * mm, bottomMargin=18 * mm)
    story = []
    base = forecast[forecast["scenario"] == "Base"].sort_values("year")

    story += [
        Spacer(1, 34 * mm),
        p("삼양식품 기업분석 및 가치평가", "KTitle"),
        p("해외 매출·생산능력·수익성을 DCF·PER·RIM으로 연결한 포트폴리오", "KSubtitle"),
        Spacer(1, 14 * mm),
        table(
            [
                ["기준일", "기업", "종목코드", "분석 범위"],
                ["2026-06-30", "삼양식품", "KRX 003230", "2021–2025A / 2026H1 / 2026E–2030E"],
            ],
            [35 * mm, 42 * mm, 35 * mm, 60 * mm],
            8.5,
        ),
        Spacer(1, 14 * mm),
        p("핵심 결론", "KH2"),
        p(f"2026H1 해외 매출 비중은 82.9%이며, Base 전망은 2026E 매출 3.10조원과 영업이익률 23.5%를 제시한다. 평가값은 DCF {valuation['DCF']:,.0f}원, PER {valuation['PER']:,.0f}원, RIM {valuation['RIM']:,.0f}원이다.", "KCallout"),
        Spacer(1, 26 * mm),
        p("교육·취업 포트폴리오용 자료이며 투자 권유가 아니다.", "KSmall"),
        PageBreak(),
    ]

    story += [
        p("1. 한눈에 보는 결론", "KH1"),
        table(
            [
                ["항목", "결과", "해석"],
                ["2025 매출 / 영업이익률", "2.352조원 / 22.3%", "고성장과 마진 상승이 함께 나타남"],
                ["2026H1 매출 / 영업이익률", "1.485조원 / 23.8%", "증설 가동과 해외 비중 확대"],
                ["2026H1 해외 매출 비중", "82.9%", "전망의 핵심 변수가 해외 수요임"],
                ["2026E Base 매출 / EPS", "3.10조원 / 76,132원", "상반기 성장률보다 완만한 하반기 가정"],
                ["DCF / PER / RIM", f"{valuation['DCF']/1e6:.2f} / {valuation['PER']/1e6:.2f} / {valuation['RIM']/1e6:.2f}백만원", "평가법별 역할을 구분해 해석"],
            ],
            [47 * mm, 45 * mm, 78 * mm],
        ),
        Spacer(1, 8 * mm),
        Image(str(figures[2]), width=165 * mm, height=93 * mm),
        p(f"2026-09-18 종가 1,300,000원 대비 DCF는 {valuation['DCF']/REFERENCE_PRICE-1:.1%}, PER은 {valuation['PER']/REFERENCE_PRICE-1:.1%}, RIM은 {valuation['RIM']/REFERENCE_PRICE-1:.1%}의 차이를 보인다.", "KSmall"),
        PageBreak(),
    ]

    story += [
        p("2. 기업과 분석 질문", "KH1"),
        p("삼양식품의 주력 제품은 불닭 브랜드를 포함한 면류다. 2026H1에는 해외 매출이 전체의 82.9%를 차지해 국내 식품기업 평균보다 해외 수요와 환율, 지역별 유통 성과의 영향이 크다."),
        p("분석 질문", "KH2"),
        p("① 해외 성장률이 둔화돼도 높은 수익성이 유지되는가? ② 밀양 2공장과 중국 자싱 공장이 매출로 연결되는가? ③ 현재 가격은 어느 수준의 성장과 ROE를 반영하는가?"),
        p("전망 방법", "KH2"),
        p("매출은 국내와 해외를 구분하고, 해외는 미국·중국·유럽·기타 지역의 수요를 살핀다. 마진은 제품·지역 믹스, 생산 효율, 원재료·물류·환율, 마케팅 비용으로 설명한다. 2027년에는 중국 현지 생산의 시작을 별도 전제로 둔다."),
        p("자료 사용 원칙", "KH2"),
        p("DART와 회사 IR을 우선 사용했다. 기사·증권사 자료는 지역별 판매와 비교기업 컨센서스 확인에 한정했으며, 공시 대사가 필요한 값은 가정표에서 구분했다."),
        PageBreak(),
    ]

    hist_rows = [["연도", "매출(십억원)", "영업이익(십억원)", "영업이익률", "순이익(십억원)"]]
    for row in hist.itertuples(index=False):
        hist_rows.append([str(row.year), f"{row.revenue_krw_mn/1000:,.1f}", f"{row.operating_profit_krw_mn/1000:,.1f}", f"{row.operating_profit_krw_mn/row.revenue_krw_mn:.1%}", f"{row.net_income_krw_mn/1000:,.1f}"])
    story += [
        p("3. 2021–2025 실적", "KH1"),
        Image(str(figures[0]), width=165 * mm, height=88 * mm),
        Spacer(1, 4 * mm),
        table(hist_rows, [25 * mm, 38 * mm, 40 * mm, 33 * mm, 36 * mm]),
        Spacer(1, 5 * mm),
        p("2021–2025년 매출은 6,420억원에서 2조 3,518억원으로 증가했다. 영업이익률은 10.2%에서 22.3%로 상승했다. 2024년 이후 성장률과 수익성이 함께 높아졌으므로, 전망에서는 해외 판매와 증설의 지속 여부를 따로 확인해야 한다."),
        PageBreak(),
    ]

    story += [
        p("4. 2026년 상반기 기준점", "KH1"),
        table(
            [
                ["항목", "2026H1", "비교 기준 / 설명"],
                ["매출", "1조 4,847억원", "전년 동기 대비 +37.2%"],
                ["영업이익", "3,533억원", "영업이익률 23.8%"],
                ["순이익", "2,821억원", "RIM과 EPS 전망의 기준"],
                ["해외 매출", "1조 2,308억원", "전체 매출의 82.9%"],
                ["자본총계", "1조 5,483억원", "RIM 기초 자기자본"],
                ["상장주식수", "7,533,015주", "주당가치 계산 기준"],
                ["순현금", "763억원", "DCF equity bridge"],
            ],
            [48 * mm, 45 * mm, 75 * mm],
        ),
        Spacer(1, 8 * mm),
        p("상반기 실적은 2025년의 높은 성장 이후에도 해외 판매와 20%대 영업이익률이 유지됐음을 보여준다. 다만 2026E Base는 상반기 성장률을 그대로 연장하지 않고, 하반기 전년 대비 성장률이 약 27%로 낮아지는 것으로 설정했다.", "KCallout"),
        p("자본총계 1조 5,483억원과 7,533,015주는 2026H1 기준 RIM 입력값으로 사용했다. 비지배지분이 유의미하게 달라질 경우 지배주주지분으로 교체해야 한다.", "KSmall"),
        PageBreak(),
    ]

    scenario_rows = [["구분", "2026E", "2027E", "2028E", "2029E", "2030E"]]
    for scenario in ["Bear", "Base", "Bull"]:
        g = forecast[forecast["scenario"] == scenario].sort_values("year")
        scenario_rows.append([f"{scenario} 매출"] + [f"{v:,.0f}" for v in g["revenue_krw_bn"]])
    scenario_rows.append(["Base 영업이익률"] + [f"{v:.1%}" for v in base["operating_margin"]])
    scenario_rows.append(["Base FCFF"] + [f"{v:,.1f}" for v in base["fcff_krw_bn"]])
    story += [
        p("5. 2026E–2030E 전망", "KH1"),
        Image(str(figures[1]), width=165 * mm, height=88 * mm),
        Spacer(1, 4 * mm),
        table(scenario_rows, [35 * mm] + [27 * mm] * 5),
        Spacer(1, 5 * mm),
        p("Base 매출 성장률은 31.8%, 20.0%, 16.0%, 12.0%, 9.0%로 점차 낮아진다. 영업이익률은 23.5~24.0% 범위로 두어 최근 마진이 계속 상승한다는 가정을 피했다. CAPEX는 증설기 이후 감소하고, 유지보수 CAPEX는 자산 증가에 맞춰 늘어난다."),
        PageBreak(),
    ]

    story += [
        p("6. DCF", "KH1"),
        p("FCFF는 NOPAT + D&A − CAPEX − ΔNWC로 계산했다. 2026E–2030E Base FCFF를 mid-year 방식으로 할인하고, 2030E 이후에는 영구성장모형을 적용했다."),
        table(
            [
                ["가정", "값", "설명"],
                ["WACC", "8.5%", "모델 독립 가정"],
                ["영구성장률", "2.5%", "장기 명목 성장률"],
                ["순현금", "763억원", "2026H1 현금·단기금융자산−차입금−리스"],
                ["주당가치", f"{valuation['DCF']:,.0f}원", f"기준 종가 대비 {valuation['DCF']/REFERENCE_PRICE-1:.1%}"],
            ],
            [55 * mm, 38 * mm, 76 * mm],
        ),
        Spacer(1, 8 * mm),
        p("DCF 민감도", "KH2"),
        table(
            [
                ["WACC / g", "2.0%", "2.5%", "3.0%"],
                ["7.5%", "1,811,710", "1,966,852", "2,156,470"],
                ["8.0%", "1,656,901", "1,784,104", "1,936,799"],
                ["8.5%", "1,525,990", f"{valuation['DCF']:,.0f}", "1,757,060"],
                ["9.0%", "1,413,853", "1,503,181", "1,607,397"],
                ["9.5%", "1,316,733", "1,392,920", "1,480,826"],
            ],
            [42 * mm] * 4,
        ),
        Spacer(1, 6 * mm),
        p("기업가치의 상당 부분이 말기가치에서 나오므로 WACC와 영구성장률의 작은 변화가 결과를 크게 바꾼다. 기준값 한 개보다 민감도 범위를 함께 제시한다.", "KSmall"),
        PageBreak(),
    ]

    story += [
        p("7. PER 상대가치", "KH1"),
        p("PER은 2027E EPS 92,839원을 사용했다. 비교기업은 라면 중심 사업의 농심과 해외 브랜드 식품기업 오리온을 1차 패널로 두었다. 2027E 두 기업 PER 중앙값 9.815배에 관찰된 삼양식품 프리미엄 약 50%를 적용해 Base 14.72배를 산출했다."),
        table(
            [
                ["구분", "적용 PER", "주당가치", "기준 종가 대비"],
                ["Bear", "13.25배", "1,230,117원", "-5.4%"],
                ["Base", "14.72배", f"{valuation['PER']:,.0f}원", f"{valuation['PER']/REFERENCE_PRICE-1:.1%}"],
                ["Bull", "16.19배", "1,503,063원", "+15.6%"],
            ],
            [38 * mm, 35 * mm, 48 * mm, 45 * mm],
        ),
        Spacer(1, 8 * mm),
        p("현재 시장에서 관찰되는 프리미엄을 목표 배수에 쓰면 순환논리가 생길 수 있다. 따라서 PER은 DCF의 독립 대체값이 아니라 비교기업 대비 가격 수준을 점검하는 보조 결과로 사용한다.", "KCallout"),
        PageBreak(),
    ]

    story += [
        p("8. RIM", "KH1"),
        p("RIM은 기초 자기자본에 미래 잔여이익의 현재가치를 더한다. 잔여이익은 순이익에서 자기자본비용을 뺀 값이다. 기존 템플릿의 ‘ROE×자기자본÷요구수익률’ 계산을 교체하고, 연도별 장부가치와 배당을 이어서 계산했다."),
        table(
            [
                ["가정", "값", "설명"],
                ["기초 자기자본", "1조 5,483억원", "2026H1 공시 기준"],
                ["자기자본비용", "9.0%", "모델 독립 가정"],
                ["배당성향", "15.0%", "순이익에서 배당 차감"],
                ["장기 ROE", "14.0%", "고성장 이후 정상화"],
                ["영구성장률", "2.5%", "장기 명목 성장률"],
                ["주당가치", f"{valuation['RIM']:,.0f}원", f"기준 종가 대비 {valuation['RIM']/REFERENCE_PRICE-1:.1%}"],
            ],
            [55 * mm, 38 * mm, 76 * mm],
        ),
        Spacer(1, 8 * mm),
        p("RIM이 DCF와 PER보다 낮은 이유", "KH2"),
        p("2026H1 장부가치에 비해 시장가치가 높은 상황에서 장기 ROE를 14%로 정상화하면 초과이익의 지속기간이 짧아진다. 현재 주가를 설명하려면 더 높은 장기 ROE, 낮은 자기자본비용, 긴 초과이익 지속기간 중 하나가 필요하다."),
        PageBreak(),
    ]

    story += [
        p("9. 위험과 확인 지표", "KH1"),
        table(
            [
                ["위험", "확인 지표", "평가 영향"],
                ["해외 성장 둔화", "미국·중국·유럽 매출 성장률", "매출·PER 하락"],
                ["증설 가동 지연", "밀양 2공장·자싱 공장 가동률", "매출·FCFF 하락"],
                ["원재료·물류비 상승", "원재료 단가·운임·매출총이익률", "마진 하락"],
                ["환율 반전", "원/달러·원/위안 환율", "수출 환산액과 마진 영향"],
                ["마케팅비 증가", "SG&A와 해외 채널 비용", "영업이익률 하락"],
                ["프리미엄 축소", "비교기업 PER·성장률 차이", "상대가치 하락"],
            ],
            [42 * mm, 63 * mm, 62 * mm],
        ),
        Spacer(1, 8 * mm),
        p("분기별 업데이트 순서", "KH2"),
        p("① 공시 실적을 전망과 비교한다. ② 지역별 매출과 공장 가동률을 갱신한다. ③ NWC·CAPEX·D&A를 다시 연결한다. ④ 비교기업 컨센서스 날짜를 맞춘다. ⑤ 세 평가법을 같은 주식수와 기준일로 재계산한다."),
        PageBreak(),
    ]

    story += [
        p("10. 재현 파일과 계산 규칙", "KH1"),
        table(
            [
                ["파일", "용도"],
                ["model/DCF_model.xlsx", "DCF 입력·전망·민감도와 주당가치"],
                ["model/RIM_model.xlsx", "장부가치 롤포워드·잔여이익·민감도"],
                ["model/PER_template.xlsx", "기존 PER 템플릿"],
                ["model/valuation_assumptions.csv", "세 평가법 공통 가정"],
                ["model/valuation_results.csv", "평가 결과 요약"],
                ["src/valuation.py", "DCF·PER·RIM 재계산"],
                ["data/processed/forecast_scenarios_2026_2030.csv", "Base·Bull·Bear 전망"],
            ],
            [68 * mm, 100 * mm],
        ),
        Spacer(1, 8 * mm),
        p("검증 규칙", "KH2"),
        p("주당가치 계산에는 7,533,015주를 공통으로 사용한다. 금액 단위는 모델에서 십억원, 표시에서 조원·억원으로 변환한다. 공시값·전망값·가정값을 구분한다. Excel 결과와 Python 결과는 같은 입력표에서 맞춘다."),
        p("현재 버전", "KH2"),
        p("2026H1 공시를 기준으로 한 v1.0이다. 2026년 연간 실적 발표 뒤 매출, 자본, 순현금, 비교기업 컨센서스를 갱신해야 한다."),
        PageBreak(),
    ]

    sources = [
        ("[1] 삼양식품 IR 재무정보", "https://www.samyangfoods.com/kor/ir/finance.do"),
        ("[2] 삼양식품 IR 자료실", "https://www.samyangfoods.com/kor/ir/list.do"),
        ("[3] DART 반기보고서, 접수번호 20260814003053", "https://dart.fss.or.kr/dsaf001/main.do?rcpNo=20260814003053"),
        ("[4] KRX KIND", "https://kind.krx.co.kr/"),
        ("[5] FnGuide CompanyGuide, 발행주식수·시장지표 확인", "https://kwcomp.fnguide.com/CompanyInfo/Snapshot?cmp_cd=003230"),
        ("[6] 2026-09-18 확정 종가 및 공시 재무 요약", "https://aikstockdata.com/s/003230"),
        ("[7] PER 비교기업 자료 기록", "research/PER_PREMIUM_FRAMEWORK.md"),
    ]
    story += [p("11. 출처와 면책", "KH1")]
    for label, url in sources:
        if url.startswith("http"):
            story.append(p(f'{label}<br/><link href="{url}" color="#2F6BFF">{url}</link>', "KSource"))
        else:
            story.append(p(f"{label}: {url}", "KSource"))
    story += [
        Spacer(1, 8 * mm),
        p("면책", "KH2"),
        p("이 보고서는 교육·취업 포트폴리오 목적으로 작성했다. 전망과 가치평가는 작성자의 가정이며, 증권 매수·매도 권유가 아니다. 시장가격과 컨센서스는 기준일 이후 달라질 수 있다."),
    ]

    doc.build(story, onFirstPage=page_number, onLaterPages=page_number)
    return path


def main() -> None:
    MODEL.mkdir(exist_ok=True)
    FIGURES.mkdir(exist_ok=True)
    REPORT.mkdir(exist_ok=True)
    hist = pd.read_csv(DATA / "financials_annual.csv", encoding="utf-8-sig")
    forecast = pd.read_csv(DATA / "forecast_scenarios_2026_2030.csv")
    base = forecast[forecast["scenario"] == "Base"].sort_values("year")
    dcf_equity, _, _ = dcf_value(base["fcff_krw_bn"].tolist(), WACC, TERMINAL_GROWTH)
    rim_equity, _ = rim_value(base["net_income_krw_bn"].tolist(), COE, TERMINAL_ROE)
    valuation = {
        "DCF": round(dcf_equity * 1000 / SHARES_M),
        "PER": round(float(base.loc[base["year"] == 2027, "eps_krw"].iloc[0]) * PER_MULTIPLE),
        "RIM": round(rim_equity * 1000 / SHARES_M),
    }
    dcf_path = build_dcf_workbook(base)
    rim_path = build_rim_workbook(base)
    figures = build_figures(hist, forecast, valuation)
    pdf_path = build_pdf(hist, forecast, valuation, figures)
    print({"DCF": str(dcf_path), "RIM": str(rim_path), "PDF": str(pdf_path), **valuation})


if __name__ == "__main__":
    main()
